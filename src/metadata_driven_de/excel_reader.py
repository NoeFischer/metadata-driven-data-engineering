"""Read pipeline metadata from Excel workbooks.

Expected Excel layout
---------------------
Sheet: **tables**
    Each row is one pipeline / table. Columns:
    pipeline_name, description, source_path, source_format, header, delimiter,
    target_catalog, target_schema, target_table, write_mode, partition_by, merge_keys

Sheet: **columns**
    Each row is one column mapping. Columns:
    pipeline_name, source_name, target_name, data_type, nullable, description, transform

The two sheets are joined on ``pipeline_name``.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _parse_list_field(value: Any) -> list[str]:
    """Turn a comma-separated cell value into a list of strings."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return []
    return [v.strip() for v in str(value).split(",") if v.strip()]


def _parse_bool(value: Any, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "yes")


def _sheet_to_dicts(ws) -> list[dict[str, Any]]:
    """Convert an openpyxl worksheet into a list of row-dicts using header row."""
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h).strip().lower() for h in rows[0]]
    return [
        {h: cell for h, cell in zip(headers, row)}
        for row in rows[1:]
        if any(cell is not None for cell in row)
    ]


def _build_read_options(row: dict[str, Any]) -> dict[str, str]:
    """Extract any column starting with ``opt_`` as a Spark read option."""
    prefix = "opt_"
    return {
        k[len(prefix):]: str(v)
        for k, v in row.items()
        if k.startswith(prefix) and v is not None
    }


def read_excel(path: str | Path) -> list[dict[str, Any]]:
    """Parse an Excel workbook into a list of raw pipeline dicts.

    Returns a list of dicts ready to be passed into
    ``IngestionPipeline.model_validate()``.
    """
    wb = load_workbook(Path(path), data_only=True)

    sheet_names_lower = {s.lower(): s for s in wb.sheetnames}
    if "tables" not in sheet_names_lower:
        raise ValueError(f"Workbook must contain a 'tables' sheet. Found: {wb.sheetnames}")
    if "columns" not in sheet_names_lower:
        raise ValueError(f"Workbook must contain a 'columns' sheet. Found: {wb.sheetnames}")

    table_rows = _sheet_to_dicts(wb[sheet_names_lower["tables"]])
    column_rows = _sheet_to_dicts(wb[sheet_names_lower["columns"]])

    # Group columns by pipeline_name
    cols_by_pipeline: dict[str, list[dict]] = {}
    for cr in column_rows:
        pname = str(cr.get("pipeline_name", "")).strip()
        if not pname:
            continue
        cols_by_pipeline.setdefault(pname, []).append(
            {
                "source_name": str(cr.get("source_name", "")),
                "target_name": str(cr.get("target_name", "")),
                "data_type": str(cr.get("data_type", "string")).strip().lower(),
                "nullable": _parse_bool(cr.get("nullable"), default=True),
                "description": str(cr.get("description") or ""),
                "transform": str(cr.get("transform") or ""),
            }
        )

    pipelines: list[dict[str, Any]] = []
    for row in table_rows:
        pname = str(row.get("pipeline_name", "")).strip()
        if not pname:
            continue

        pipeline = {
            "pipeline_name": pname,
            "description": str(row.get("description") or ""),
            "source": {
                "path": str(row.get("source_path", "")),
                "format": str(row.get("source_format", "csv")).strip().lower(),
                "header": _parse_bool(row.get("header"), default=True),
                "delimiter": str(row.get("delimiter") or ","),
                "read_options": _build_read_options(row),
            },
            "target": {
                "catalog": str(row.get("target_catalog", "")),
                "schema": str(row.get("target_schema", "")),
                "table": str(row.get("target_table", "")),
                "write_mode": str(row.get("write_mode") or "append").strip().lower(),
                "partition_by": _parse_list_field(row.get("partition_by")),
                "merge_keys": _parse_list_field(row.get("merge_keys")),
            },
            "columns": cols_by_pipeline.get(pname, []),
        }
        pipelines.append(pipeline)

    return pipelines
